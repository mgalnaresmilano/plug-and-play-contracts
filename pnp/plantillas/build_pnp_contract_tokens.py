"""Build PlugAndPlay_Contract_Tokens.docx, the English Ecosystem Plus Membership
Agreement template.

Starts from the Plug and Play demo contract, keeps its styling, and replaces the
hard-coded Spanish values with Context Service tokens from PNP_ContractDocGenContext
(Contract -> Quote -> LineItems). The wording follows Plug and Play's real
Ecosystem Plus Membership Agreement.

The clause content controls are rebuilt from the clause library defined in
clausulas_pnp.py. Server-side generation does not re-read the library, so the
clause text has to be written into the document; the control still points at the
library record so the Word add-in can track and swap it.

Usage: python3 build_pnp_contract_tokens.py [--org plug-and-play-trailsignup]
"""

import argparse
import copy
import hashlib
import json
import shutil
import zipfile
from pathlib import Path

import docx
from docx.oxml.ns import qn

from clausulas_pnp import CLAUSES, Org

HERE = Path(__file__).parent
SOURCE = HERE / "original" / "PlugAndPlay_Contract_Template.docx"
BANNER = HERE / "original" / "banner_en.png"
OUTPUT = HERE / "PlugAndPlay_Contract_Tokens.docx"

PNP_ADDRESS = [
    "Plug & Play, LLC",
    "440 North Wolfe Road",
    "94085 Sunnyvale, CA (United States)",
]


def set_text(paragraph, text):
    """Replace a paragraph's text keeping the formatting of its first run."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def set_element_text(p_element, text):
    """Same as set_text but on a raw w:p element (used inside content controls)."""
    runs = p_element.findall(qn("w:r"))
    for run in runs[1:]:
        p_element.remove(run)
    t = runs[0].find(qn("w:t"))
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def insert_after(element, template_paragraph, text):
    """Insert a copy of template_paragraph after element and set its text."""
    new_p = copy.deepcopy(template_paragraph._p)
    element.addnext(new_p)
    paragraph = docx.text.paragraph.Paragraph(new_p, template_paragraph._parent)
    set_text(paragraph, text)
    return paragraph


def body_paragraph(document, startswith):
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(startswith):
            return paragraph
    raise ValueError(f"Paragraph starting with {startswith!r} not found")


def clause_ids(org):
    names = "', '".join(c["name"] for c in CLAUSES)
    found = org.query("SELECT Id, Name, DocumentClauseSet.Name FROM DocumentClause "
                      f"WHERE Name IN ('{names}') AND Status = 'Active'")
    return {r["Name"]: (r["Id"], r["DocumentClauseSet"]["Name"]) for r in found}


def clause_guid(index, clause_id):
    return f"pnp-clm-{index}-{clause_id[-8:]}"


def link_clauses(org, ids, template_name):
    """Register each clause control on the template as DocumentAuthoredContent.

    The Contracts Connector for Word resolves every clause control through this
    record (matched by the guid in the control's tag); without it the add-in
    panel fails to load on documents generated from the template.
    """
    found = org.query(f"SELECT Id FROM DocumentTemplate WHERE Name = '{template_name}' "
                      "AND Status = 'Active' ORDER BY CreatedDate DESC LIMIT 1")
    if not found:
        print(f"Template {template_name} not published yet: run publicar_plantilla.py and build again")
        return
    template_id = found[0]["Id"]
    existing = {r["DocAuthoredContentIdentifier"]: r["Id"] for r in org.query(
        "SELECT Id, DocAuthoredContentIdentifier FROM DocumentAuthoredContent "
        f"WHERE ReferenceObjectId = '{template_id}'")}
    wanted = {}
    for index, clause in enumerate(CLAUSES, start=1):
        clause_id = ids[clause["name"]][0]
        wanted[clause_guid(index, clause_id)] = clause_id

    for guid, clause_id in wanted.items():
        if guid in existing:
            continue
        org.call("POST", "sobjects/DocumentAuthoredContent", {
            "ReferenceObjectId": template_id,
            "ReferenceObjectRecordType": "DocumentTemplate",
            "ContentType": "StandardDocumentClause",
            "StandardContentObjectType": "DocumentClause",
            "StandardContentObjectId": clause_id,
            "DocAuthoredContentIdentifier": guid,
            "ContentGenerationSource": "StandardContent",
            "IsReviewed": False,
            "IsLibraryAdditionRequested": False,
        })
    for guid, record_id in existing.items():
        if guid not in wanted:
            org.call("DELETE", f"sobjects/DocumentAuthoredContent/{record_id}")
    print(f"Clause controls registered on {template_name}: {len(wanted)}")


def rebuild_clauses(document, ids):
    """Swap the template's clause controls for the English tokenized clauses."""
    body = document.element.body
    controls = body.findall(qn("w:sdt"))
    prototype = controls[0]
    proto_content = prototype.find(qn("w:sdtContent"))
    proto_paragraphs = proto_content.findall(qn("w:p"))
    heading_proto, body_proto = proto_paragraphs[0], proto_paragraphs[1]

    anchor = prototype
    for index, clause in enumerate(CLAUSES, start=1):
        clause_id, set_name = ids[clause["name"]]
        block = copy.deepcopy(prototype)
        content = block.find(qn("w:sdtContent"))
        for paragraph in content.findall(qn("w:p")):
            content.remove(paragraph)

        heading = copy.deepcopy(heading_proto)
        set_element_text(heading, clause["heading"])
        content.append(heading)
        for text in clause["body"]:
            paragraph = copy.deepcopy(body_proto)
            set_element_text(paragraph, text)
            content.append(paragraph)

        text_blob = "\n".join([clause["heading"], *clause["body"]])
        tag = block.find(qn("w:sdtPr")).find(qn("w:tag"))
        tag.set(qn("w:val"), json.dumps({
            "entityId": clause_id,
            "type": "Clause",
            "title": clause["name"],
            "language": "English",
            "category": set_name,
            "generateSource": "Library",
            "controlId": 100000000 + int(hashlib.sha256(clause_id.encode()).hexdigest(), 16) % 900000000,
            "guid": clause_guid(index, clause_id),
            "checksum": hashlib.sha256(text_blob.encode("utf-8")).hexdigest(),
        }, separators=(",", ":"), ensure_ascii=False))

        anchor.addnext(block)
        anchor = block

    for control in controls:
        body.remove(control)


def swap_banner(path):
    """Replace the cover bitmap, whose title is baked into the artwork."""
    temp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = BANNER.read_bytes() if item.filename == "word/media/image1.png" \
                else source.read(item.filename)
            target.writestr(item, data)
    shutil.move(temp, path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--org", default="plug-and-play-trailsignup")
    args = parser.parse_args()

    org = Org(args.org)
    ids = clause_ids(org)
    missing = [c["name"] for c in CLAUSES if c["name"] not in ids]
    if missing:
        raise SystemExit(f"Clauses missing in the org, run clausulas_pnp.py first: {missing}")

    document = docx.Document(SOURCE)
    header, products, totals, discount_box, conditions, signatures = document.tables

    title = document.paragraphs[1]
    set_text(title, "Ecosystem Plus Membership Agreement")

    intro = body_paragraph(document, "Membresía corporativa")
    set_text(
        intro,
        "This Ecosystem Plus Membership Agreement (\u201cAgreement\u201d) is made effective as of "
        "{{DT_StartDate}} (the \u201cEffective Date\u201d), by and between {{AccountName}}, whose "
        "principal office is located at {{DT_BillingStreet}}, {{DT_BillingPostalCode}} "
        "{{DT_BillingCity}} ({{DT_BillingCountry}}) (hereinafter \u201cClient\u201d), and "
        "Plug & Play, LLC (also doing business as \u201cPlug and Play Tech Center\u201d), a California "
        "limited liability company, whose principal office is located at 440 North Wolfe Road, "
        "Sunnyvale, CA 94085, United States (\u201cPlug and Play\u201d). The Parties record this "
        "Agreement in Salesforce as contract {{ContractNumber}}.",
    )

    # Header: Plug and Play | client | term | agreement
    provider, client, validity, contract = header.rows[0].cells
    set_text(provider.paragraphs[0], "PLUG AND PLAY")
    provider_value = provider.paragraphs[1]
    set_text(provider_value, PNP_ADDRESS[0])
    last = provider_value._p
    for line in PNP_ADDRESS[1:]:
        last = insert_after(last, provider_value, line)._p

    set_text(client.paragraphs[0], "CLIENT")
    client_value = client.paragraphs[1]
    set_text(client_value, "{{AccountName}}")
    last = client_value._p
    for line in (
        "{{AccountBillingStreet}}",
        "{{AccountBillingPostalCode}} {{AccountBillingCity}} ({{AccountBillingCountry}})",
        "{{AccountPhone}}",
    ):
        last = insert_after(last, client_value, line)._p

    set_text(validity.paragraphs[0], "TERM")
    set_text(validity.paragraphs[1], "Effective Date - {{ContractEndDate}}")
    insert_after(validity.paragraphs[1]._p, validity.paragraphs[1],
                 "Status: {{ContractStatus}}")

    set_text(contract.paragraphs[0], "AGREEMENT")
    set_text(contract.paragraphs[1], "No. {{ContractNumber}}")

    # Exhibit A opens the Quote block: quote and opportunity tokens only resolve
    # between {{#Quote}} and {{/Quote}}.
    economics = body_paragraph(document, "Condiciones económicas")
    set_text(economics, "Exhibit A: Fee Schedule")
    quote_line = insert_after(
        economics._p, intro,
        "{{#Quote}}Proposal {{QuoteNumber}} · {{QuoteName}} · Status: {{QuoteStatus}} · "
        "Issued on {{QuoteCreatedDate}} · Valid until {{QuoteExpirationDate}}",
    )
    insert_after(
        quote_line._p, intro,
        "Opportunity: {{OpportunityName}} · Stage: {{OpportunityStage}} · Type: {{OpportunityType}}",
    )

    # Fee schedule: row 1 is the template repeated for every quote line.
    head = products.rows[0].cells
    for cell, label in zip(head[1:], ("DESCRIPTION OF WORK", "QTY", "LIST PRICE (USD)",
                                      "DISC. %", "AMOUNT (USD)")):
        set_text(cell.paragraphs[0], label)
    for row in list(products.rows)[2:]:
        row._tr.getparent().remove(row._tr)
    marker, product, quantity, list_price, line_discount, line_total = products.rows[1].cells
    set_text(marker.paragraphs[0], "{{#LineItems}}")
    set_text(product.paragraphs[0], "{{ProductName}}")
    set_text(product.paragraphs[1], "SKU {{ProductCode}}")
    set_text(product.paragraphs[2], "{{ProductDescription}}")
    set_text(quantity.paragraphs[0], "{{Quantity}}")
    set_text(list_price.paragraphs[0], "{{ListPrice}}")
    set_text(line_discount.paragraphs[0], "{{LineDiscount}}")
    set_text(line_total.paragraphs[0], "{{LineTotal}}{{/LineItems}}")

    # QuoteDiscount is an unrounded percentage (10.79136691), so the totals show
    # amounts only and the per-line DISCOUNT column carries the percentages.
    subtotal_row, net_row, value_row = totals.rows
    set_text(subtotal_row.cells[0].paragraphs[0], "Subtotal (list price)")
    set_text(subtotal_row.cells[1].paragraphs[0], "USD {{QuoteSubtotal}}")
    net_row._tr.getparent().remove(net_row._tr)
    set_text(value_row.cells[0].paragraphs[0], "Total agreement value")
    set_text(value_row.cells[1].paragraphs[0], "USD {{QuoteGrandTotal}}")

    box = discount_box.rows[0].cells[0]
    set_text(box.paragraphs[0], "Approved discount")
    set_text(
        box.paragraphs[1],
        "The discounts applied to proposal {{QuoteNumber}} were approved under the Plug and Play "
        "approval matrix and are conditional on Client meeting the full membership term set out in "
        "the Program Terms. The resulting agreement value is USD {{QuoteGrandTotal}}.",
    )
    # PnPDiscountApprover lives on the Contract node, so it goes after the section closes.
    insert_after(discount_box._tbl, intro,
                 "{{/Quote}}Discount approved by {{PnPDiscountApprover}}.")

    # Program terms: everything here comes from the contract. DT_ tokens are the
    # Data True-Up fields: each may appear only once, outside the Quote section.
    set_text(body_paragraph(document, "Condiciones del programa"), "Program Terms")
    program_terms = [
        ("Membership term", "{{DT_ContractTerm}} months"),
        ("Kickoff date", "{{DT_PnPKickoffDate}}"),
        ("Workshops included", "{{DT_PnPWorkshopsIncluded}} innovation workshops"),
        ("Marketing hours included", "{{DT_PnPMarketingHours}} hours of marketing support"),
        ("Relationship manager", "{{DT_PnPProgramManager}}"),
        ("Client executive sponsor", "{{DT_PnPExecutiveSponsor}}"),
        ("Source proposal", "{{PnPSourceQuoteNumber}}"),
        ("Renewal review", "{{DT_PnPRenewalDate}} · 60 days' written notice required"),
    ]
    for row, (label, value) in zip(conditions.rows, program_terms):
        set_text(row.cells[0].paragraphs[0], label)
        set_text(row.cells[1].paragraphs[0], value)

    set_text(body_paragraph(document, "Cláusulas contractuales"), "Terms and Conditions")
    rebuild_clauses(document, ids)

    set_text(body_paragraph(document, "Firma"), "Acknowledged, Accepted and Agreed")
    pnp_signature, client_signature = signatures.rows[0].cells
    set_text(pnp_signature.paragraphs[0], "BY PLUG & PLAY, LLC")
    set_text(pnp_signature.paragraphs[2], "Date")
    set_text(client_signature.paragraphs[0], "BY {{AccountName}}")
    set_text(client_signature.paragraphs[1], "Authorized representative of {{AccountName}}")
    set_text(client_signature.paragraphs[2], "Date")

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            if paragraph.text.startswith("Contrato confidencial"):
                set_text(paragraph,
                         "Confidential  |  Agreement {{ContractNumber}}  |  {{AccountName}}")

    document.save(OUTPUT)
    swap_banner(OUTPUT)
    print(f"Saved {OUTPUT}")
    link_clauses(org, ids, OUTPUT.stem)


if __name__ == "__main__":
    main()
